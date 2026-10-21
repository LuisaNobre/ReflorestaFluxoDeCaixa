"""Extrai atributos do Excel para os filtros de cenário do dashboard estático.

Importa VPLiquido e InvNecessario por hectare sem escolher automaticamente um fluxo.
"""

import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
ENTRADA = RAIZ / "data" / "refloresta" / "MelhorFluxoDeCaixa.xlsx"
SAIDA = RAIZ / "docs" / "data" / "filtros_cenarios.json"
CAMPOS = [
    "idRegiaoEco", "nomeRegiaoEco", "idFitofisionomia", "descFitofisionomia",
    "idRegiaoAdm", "nomeRegiaoAdm", "idFaixaTipo", "nomeFaixa",
    "idTopografia", "nomeTopografia", "idMecanizacaoNivel", "nomeMecanizacao",
    "idCombinacao", "idFluxoCaixa",
]
FILTROS = [
    ("faixa", "idFaixaTipo", "nomeFaixa"),
    ("topografia", "idTopografia", "nomeTopografia"),
    ("mecanizacao", "idMecanizacaoNivel", "nomeMecanizacao"),
]


def texto(valor, campo, linha):
    if valor is None or isinstance(valor, bool):
        raise ValueError(f"{campo} vazio ou inválido na linha {linha}")
    if isinstance(valor, (int, float)):
        if int(valor) != valor:
            raise ValueError(f"Identificador não inteiro em {campo}, linha {linha}")
        valor = str(int(valor))
    valor = str(valor).strip()
    if not valor or valor.startswith("="):
        raise ValueError(f"{campo} deve conter um valor explícito na linha {linha}")
    return valor


def main():
    with ENTRADA.open("rb") as arquivo:
        workbook = load_workbook(arquivo, read_only=True, data_only=False)
        candidatas = []
        for aba in workbook:
            cabecalhos = next(aba.iter_rows(min_row=1, max_row=1, values_only=True), ())
            nomes = [str(c).strip() if c is not None else "" for c in cabecalhos]
            if set(CAMPOS).issubset(nomes):
                candidatas.append((aba, nomes))
        if len(candidatas) != 1:
            workbook.close()
            raise ValueError("Esperada uma única aba com os cabeçalhos necessários aos filtros")
        aba, nomes = candidatas[0]
        indices = {campo: nomes.index(campo) for campo in CAMPOS + ["VPLiquido", "InvNecessario"]}
        registros = []
        for numero, linha in enumerate(aba.iter_rows(min_row=2, values_only=True), 2):
            if not any(v is not None for v in linha):
                continue
            registro = {campo: texto(linha[indices[campo]], campo, numero) for campo in CAMPOS}
            for campo in ["VPLiquido", "InvNecessario"]:
                valor = linha[indices[campo]]
                if isinstance(valor, bool) or not isinstance(valor, (int, float)) or not math.isfinite(valor):
                    raise ValueError(f"{campo} não numérico na linha {numero}")
                registro[campo] = valor
            registro["linha_excel"] = numero
            registros.append(registro)
        nome_aba = aba.title
        workbook.close()
    if not registros:
        raise ValueError("A planilha não contém registros para os filtros")
    opcoes = {}
    for chave, campo_id, campo_nome in FILTROS:
        valores = {}
        for registro in registros:
            identificador, nome = registro[campo_id], registro[campo_nome]
            if identificador in valores and valores[identificador] != nome:
                raise ValueError(f"Nomes conflitantes para {campo_id}={identificador}")
            valores[identificador] = nome
        opcoes[chave] = [{"id": k, "nome": v} for k, v in sorted(valores.items(), key=lambda item: item[1])]
    resultado = {
        "versao_schema": 1,
        "gerado_em": datetime.now(timezone.utc).isoformat(),
        "fonte": {"arquivo": ENTRADA.relative_to(RAIZ).as_posix(), "aba": nome_aba,
                  "sha256": hashlib.sha256(ENTRADA.read_bytes()).hexdigest()},
        "transformacao": "IDs como texto e atributos financeiros originais VPLiquido e InvNecessario, em reais por hectare.",
        "total_registros": len(registros),
        "opcoes": opcoes,
        "registros": registros,
    }
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    temporario = SAIDA.with_suffix(".tmp.json")
    temporario.write_text(json.dumps(resultado, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    temporario.replace(SAIDA)
    print(f"Catálogo gerado: {SAIDA}")
    print(f"{len(registros)} registros; " + "; ".join(f"{k}: {len(v)} opções" for k, v in opcoes.items()))


if __name__ == "__main__":
    main()
